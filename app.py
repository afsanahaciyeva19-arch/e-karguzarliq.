from flask import Flask, request, redirect, url_for, send_from_directory, Response, jsonify
import sqlite3
import os
import sys
import webbrowser
import threading
import mimetypes
import shutil
from datetime import datetime, date
from pathlib import Path
from werkzeug.utils import secure_filename

# ============================================================
# E-KARGÜZARLIQ
# ============================================================

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATA_DIR = os.environ.get("E_KARGUZAR_DATA_DIR", BASE_DIR)
os.makedirs(DATA_DIR, exist_ok=True)

DB = os.path.join(DATA_DIR, "karguzarliq.db")
UPLOAD_DIR = os.path.join(DATA_DIR, "uploads")

# İlk server quraşdırılmasında paketdəki mövcud məlumatları daimi diskə köçür.
if DATA_DIR != BASE_DIR:
    bundled_db = os.path.join(BASE_DIR, "karguzarliq.db")
    bundled_uploads = os.path.join(BASE_DIR, "uploads")
    if not os.path.exists(DB) and os.path.exists(bundled_db):
        shutil.copy2(bundled_db, DB)
    if not os.path.exists(UPLOAD_DIR) and os.path.isdir(bundled_uploads):
        shutil.copytree(bundled_uploads, UPLOAD_DIR)

os.makedirs(UPLOAD_DIR, exist_ok=True)

app = Flask(__name__)

# ============================================================
# 26 ƏSAS PAPKA
# ============================================================

FOLDERS = [
    ("01", "Gələn sənədlər"),
    ("02", "İmza üçün"),
    ("03", "Dərkənar / Rezolyusiya / Tapşırıq"),
    ("04", "İcrada olan"),
    ("05", "Göndəriləcək"),
    ("06", "Göndərilən sənədlər"),
    ("07", "Rektorluqdan daxil olan sənədlər"),
    ("08", "Fakültələr üzrə sənədlər"),
    ("09", "Kafedralar üzrə sənədlər"),
    ("10", "Struktur bölmələri ilə yazışmalar"),
    ("11", "Digər təşkilatlarla yazışmalar"),
    ("12", "Ərizələr və müraciətlər"),
    ("13", "Əmrlər və sərəncamlar"),
    ("14", "İclaslar və protokollar"),
    ("15", "Qərarlar və tapşırıqlar"),
    ("16", "İcra nəzarəti"),
    ("17", "Hesabatlar və məlumatlar"),
    ("18", "İş planları"),
    ("19", "Tədbirlər və tədbir planları"),
    ("20", "Komissiyalar və işçi qrupları"),
    ("21", "Tələbələrlə bağlı sənədlər"),
    ("22", "Tədris məsələləri"),
    ("23", "Müəllim və əməkdaşlarla bağlı sənədlər"),
    ("24", "Maliyyə / təsərrüfat məsələləri"),
    ("25", "Müxtəlif sənədlər"),
    ("26", "Arxiv sənədləri"),
]

FOLDER_NAMES = dict(FOLDERS)

STATUSES = [
    "Yeni",
    "Baxılır",
    "İcradadır",
    "İmza üçün",
    "Göndəriləcək",
    "İcra olunub",
    "Arxivdə",
]

ICONS = {
    "01": "📥",
    "02": "🖊️",
    "03": "📌",
    "04": "⏳",
    "05": "📤",
    "06": "✉️",
    "07": "🏛️",
    "08": "🎓",
    "09": "👥",
    "10": "🏢",
    "11": "🌐",
    "12": "👤",
    "13": "📄",
    "14": "📅",
    "15": "⚖️",
    "16": "🎯",
    "17": "📊",
    "18": "📋",
    "19": "🎉",
    "20": "👥",
    "21": "🎓",
    "22": "📘",
    "23": "👤",
    "24": "💰",
    "25": "📁",
    "26": "🗄️",
}

# ============================================================
# DATABASE
# ============================================================

def db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            reg_no TEXT UNIQUE,
            folder TEXT NOT NULL,
            doc_type TEXT,
            doc_date TEXT,
            received_date TEXT,
            subject TEXT,
            sender TEXT,
            recipient TEXT,
            executor TEXT,
            deadline TEXT,
            status TEXT,
            resolution TEXT,
            notes TEXT,
            file_name TEXT,
            created_at TEXT
        )
    """)

    conn.commit()
    conn.close()


def next_reg_no():
    conn = db()

    row = conn.execute(
        "SELECT id FROM documents ORDER BY id DESC LIMIT 1"
    ).fetchone()

    conn.close()

    number = 1 if not row else row["id"] + 1
    year = datetime.now().year

    return f"{number:04d}/{year}"


def now_text():
    return datetime.now().strftime("%d.%m.%Y %H:%M")


def format_datetime(value):
    if not value:
        return ""

    try:
        if "T" in value:
            d = datetime.fromisoformat(value)
            return d.strftime("%d.%m.%Y %H:%M")
    except Exception:
        pass

    return value


def normalize(value):
    return (value or "").strip()


# ============================================================
# HTML
# ============================================================

HTML = r"""
<!doctype html>
<html lang="az">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="theme-color" content="#062b52">
<meta name="apple-mobile-web-app-capable" content="yes">
<meta name="apple-mobile-web-app-status-bar-style" content="default">
<meta name="apple-mobile-web-app-title" content="E-Kargüzarlıq">
<link rel="manifest" href="{{ url_for('manifest') }}">

<title>{% block title %}Elektron Kargüzarlıq{% endblock %}</title>

<style>

* {
    box-sizing: border-box;
}

html,
body {
    margin: 0;
    padding: 0;
    width: 100%;
    min-height: 100%;
    font-family: Arial, Helvetica, sans-serif;
    background: #f4f7fb;
    color: #13233d;
}

.layout {
    display: flex;
    width: 100%;
    min-height: 100vh;
}

/* =========================
   SOL MENYU
   ========================= */

.sidebar {
    width: 290px;
    background: linear-gradient(180deg,#062b52,#07345f);
    color: white;
    position: fixed;
    left: 0;
    top: 0;
    bottom: 0;
    overflow-y: auto;
    padding: 14px 10px;
}

.logo {
    display: flex;
    align-items: center;
    gap: 12px;
    padding: 10px 12px 20px;
    font-size: 18px;
    font-weight: bold;
}

.logo-circle {
    width: 42px;
    height: 42px;
    min-width: 42px;
    border-radius: 50%;
    background: white;
    color: #1473e6;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 22px;
}

.home {
    background: #1174e8 !important;
    font-weight: bold;
}

.menu-title {
    font-size: 11px;
    opacity: .7;
    padding: 12px;
}

.menu-item {
    display: flex;
    align-items: center;
    gap: 8px;
    color: white;
    text-decoration: none;
    padding: 8px 12px;
    border-radius: 6px;
    font-size: 13px;
    margin: 2px 0;
}

.menu-item:hover {
    background: rgba(255,255,255,.12);
}

.menu-item.active {
    background: rgba(255,255,255,.15);
}

.folder-number {
    width: 24px;
    min-width: 24px;
    font-size: 12px;
    opacity: .85;
}

.folder-icon {
    width: 22px;
    min-width: 22px;
    text-align: center;
}

.folder-name {
    line-height: 1.25;
}

/* =========================
   ƏSAS
   ========================= */

.main {
    margin-left: 290px;
    width: calc(100% - 290px);
    min-width: 0;
    min-height: 100vh;
}

.header {
    width: 100%;
    min-height: 105px;
    background: white;
    border-bottom: 1px solid #e2e8f0;
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 18px 28px;
    gap: 20px;
}

.header h1 {
    margin: 0;
    font-size: 28px;
    color: #102a4c;
}

.header p {
    margin: 5px 0 0;
    color: #64748b;
}

.date-box {
    border: 1px solid #dbe3ee;
    padding: 12px 20px;
    border-radius: 9px;
    background: white;
    white-space: nowrap;
}

.content {
    width: 100%;
    padding: 20px 22px;
}

/* =========================
   BAŞLIQ
   ========================= */

.page-title {
    margin: 0 0 16px 0;
    color: #102a4c;
}

.page-subtitle {
    margin: -8px 0 18px 0;
    color: #64748b;
}

/* =========================
   KARTLAR
   ========================= */

.cards {
    display: grid;
    grid-template-columns: repeat(4,1fr);
    gap: 14px;
    margin-bottom: 18px;
}

.card {
    background: white;
    border: 1px solid #e0e7f0;
    border-radius: 10px;
    padding: 18px;
    min-height: 105px;
}

.card-title {
    color: #2864ad;
    font-size: 14px;
    font-weight: bold;
}

.card-number {
    font-size: 28px;
    font-weight: bold;
    margin-top: 10px;
}

.card-small {
    color: #64748b;
    margin-top: 4px;
    font-size: 12px;
}

/* =========================
   PANEL
   ========================= */

.panel,
.form-page {
    background: white;
    border: 1px solid #e0e7f0;
    border-radius: 10px;
    overflow: hidden;
}

.form-page {
    padding: 22px;
}

.toolbar {
    padding: 13px;
    border-bottom: 1px solid #e5eaf1;
    display: flex;
    gap: 8px;
    align-items: center;
    flex-wrap: wrap;
}

button,
.btn {
    border: 0;
    background: #1473e6;
    color: white;
    padding: 10px 15px;
    border-radius: 6px;
    text-decoration: none;
    cursor: pointer;
    display: inline-block;
    font-size: 13px;
}

.btn.secondary {
    background: white;
    color: #29425f;
    border: 1px solid #d7e0eb;
}

.btn.danger {
    background: #dc2626;
}

.back-btn {
    margin-bottom: 14px;
}

.search {
    margin-left: auto;
    padding: 10px;
    width: 260px;
    max-width: 100%;
    border: 1px solid #d7e0eb;
    border-radius: 6px;
}

select,
input,
textarea {
    width: 100%;
    padding: 10px;
    border: 1px solid #d6dfeb;
    border-radius: 6px;
    font-family: inherit;
    font-size: 13px;
}

textarea {
    resize: vertical;
}

/* =========================
   CƏDVƏL
   ========================= */

.table-wrap {
    overflow-x: auto;
}

table {
    width: 100%;
    border-collapse: collapse;
}

th {
    background: #f7f9fc;
    color: #344a64;
    font-size: 12px;
    text-align: left;
    padding: 11px;
    border-bottom: 1px solid #e2e8f0;
}

td {
    padding: 10px;
    border-bottom: 1px solid #edf1f5;
    font-size: 12px;
}

tr:hover {
    background: #f8fbff;
}

.status {
    padding: 5px 8px;
    border-radius: 18px;
    font-size: 11px;
    background: #e8f1ff;
    color: #1463c4;
    white-space: nowrap;
}

.status-danger {
    background: #fee2e2;
    color: #b91c1c;
}

.status-ok {
    background: #dcfce7;
    color: #15803d;
}

/* =========================
   FORM
   ========================= */

.form-grid {
    display: grid;
    grid-template-columns: repeat(2,1fr);
    gap: 15px;
}

.form-group label {
    display: block;
    margin-bottom: 5px;
    font-weight: bold;
    font-size: 12px;
}

.full {
    grid-column: 1/-1;
}

.actions {
    margin-top: 18px;
    display: flex;
    gap: 8px;
}

/* =========================
   DETAL
   ========================= */

.detail-grid {
    display: grid;
    grid-template-columns: 180px 1fr;
    gap: 9px 14px;
}

.detail-grid strong {
    color: #52647b;
}

.detail-grid span {
    overflow-wrap: anywhere;
}

/* =========================
   PAPKALAR
   ========================= */

.folder-list {
    display: grid;
    grid-template-columns: repeat(4,1fr);
    gap: 12px;
}

.folder-card {
    background: white;
    border: 1px solid #dfe7f0;
    padding: 15px;
    border-radius: 8px;
    text-decoration: none;
    color: #173a64;
}

.folder-card:hover {
    border-color: #1976e8;
}

/* =========================
   BOŞ
   ========================= */

.empty {
    text-align: center;
    padding: 70px 20px;
    color: #64748b;
}

.empty-icon {
    font-size: 48px;
    margin-bottom: 10px;
}

/* =========================
   RESPONSIVE
   ========================= */

@media(max-width:1100px) {
    .cards {
        grid-template-columns: repeat(2,1fr);
    }

    .folder-list {
        grid-template-columns: repeat(2,1fr);
    }
}

@media(max-width:800px) {
    .sidebar {
        width: 220px;
    }

    .main {
        margin-left: 220px;
        width: calc(100% - 220px);
    }

    .form-grid {
        grid-template-columns: 1fr;
    }

    .full {
        grid-column: auto;
    }
}


/* MOBIL / PWA */
.mobile-menu-btn{display:none;position:fixed;top:12px;left:12px;z-index:1201;width:44px;height:44px;border-radius:10px;padding:0;font-size:23px;box-shadow:0 4px 14px rgba(0,0,0,.18)}
.sidebar-overlay{display:none;position:fixed;inset:0;background:rgba(2,14,29,.48);z-index:1190}
@media(max-width:800px){body{overflow-x:hidden}.mobile-menu-btn{display:block}.sidebar{width:min(86vw,320px);z-index:1200;transform:translateX(-105%);transition:transform .22s ease;box-shadow:8px 0 28px rgba(0,0,0,.20);padding-top:68px}body.menu-open .sidebar{transform:translateX(0)}body.menu-open .sidebar-overlay{display:block}.main{margin-left:0;width:100%}.header{min-height:auto;padding:68px 16px 16px;align-items:flex-start;flex-direction:column;gap:10px}.header h1{font-size:22px}.header p{font-size:13px}.date-box{width:100%;padding:10px 12px}.content{padding:14px 10px 24px}.cards{grid-template-columns:1fr 1fr;gap:10px}.card{min-height:90px;padding:14px}.card-number{font-size:24px}.folder-list{grid-template-columns:1fr}.form-page{padding:14px}.detail-grid{grid-template-columns:1fr}.detail-grid strong{margin-top:8px}.toolbar{align-items:stretch}.search{margin-left:0;width:100%}.actions{flex-wrap:wrap}.actions .btn,.actions button{flex:1 1 140px;text-align:center}th,td{white-space:nowrap}}
@media(max-width:480px){.cards{grid-template-columns:1fr}.header h1{font-size:20px}}
</style>
</head>

<body>

<button class="mobile-menu-btn" id="mobileMenuBtn" type="button" aria-label="Menyunu aç">☰</button>
<div class="sidebar-overlay" id="sidebarOverlay"></div>

<div class="layout">

<aside class="sidebar">

<div class="logo">
    <div class="logo-circle">✦</div>
    <div>Elektron Kargüzarlıq</div>
</div>

<a class="menu-item home" href="/">
    🏠 Ana səhifə
</a>

<div class="menu-title">PROREKTORLUĞUN PAPKA SİSTEMİ</div>

{% for code,name in folders %}

<a class="menu-item {% if current_folder == code %}active{% endif %}"
   href="{{ url_for('documents', folder=code) }}">

    <span class="folder-icon">{{ icons.get(code,'📁') }}</span>
    <span class="folder-number">{{ code }}</span>
    <span class="folder-name">{{ name }}</span>

</a>

{% endfor %}

<div class="menu-title">PARAMETRLƏR</div>

<a class="menu-item" href="{{ url_for('folders_page') }}">
    ⚙️ Papkalar
</a>

</aside>


<main class="main">

<header class="header">

<div>
    <h1>ELEKTRON KARGÜZARLIQ</h1>
    <p>Prorektorluğun sənəd uçotu və icra nəzarəti sistemi</p>
</div>

<div class="date-box">
    📅 <strong>{{ now_display }}</strong>
</div>

</header>

<section class="content">

{% block content %}{% endblock %}

</section>

</main>

</div>



<script>
(function () {

    const MAIN_SCROLL_KEY =
        "ekarguzarlıq_main_page_scroll";

    function saveMainScroll() {

        try {

            sessionStorage.setItem(
                MAIN_SCROLL_KEY,
                String(window.scrollY || window.pageYOffset || 0)
            );

        } catch (e) {}
    }


    function restoreMainScroll() {

        try {

            const saved =
                sessionStorage.getItem(
                    MAIN_SCROLL_KEY
                );

            if (saved !== null) {

                const position =
                    parseInt(saved, 10) || 0;

                window.scrollTo(
                    0,
                    position
                );

            }

        } catch (e) {}
    }


    /*
       ESAS SEHIFE XETKESI HAREKET ETDIKCE
       YADDA SAXLA
    */

    window.addEventListener(
        "scroll",
        saveMainScroll,
        { passive: true }
    );


    /*
       ISTENILEN FAYLA/SENEDE KLIK EDILMEZDƏN
       EVVEL XETKESI YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest("a");

            if (!link) {
                return;
            }

            if (
                link.target === "_blank" ||
                link.hasAttribute("download")
            ) {
                return;
            }

            saveMainScroll();

        },
        true
    );


    /*
       SEHIFE GERI ACILANDA
       EVVELKI XETKESI QAYTAR
    */

    window.addEventListener(
        "pageshow",
        function () {

            restoreMainScroll();

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                150
            );

            setTimeout(
                restoreMainScroll,
                300
            );

            setTimeout(
                restoreMainScroll,
                600
            );

        }
    );


    /*
       SEHIFE YUKLENENDE
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                200
            );

        }
    );


    /*
       PROQRAM BAGLANANDA
    */

    window.addEventListener(
        "beforeunload",
        saveMainScroll
    );

})();
</script>
<script>
(function () {

    const KEY = "EKARGUZARLIQ_SIDEBAR_SCROLL";

    let lastPosition = 0;

    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    function savePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        lastPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(lastPosition)
            );

        } catch (e) {}

    }


    function restorePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        let position = lastPosition;

        try {

            const saved =
                localStorage.getItem(KEY);

            if (saved !== null) {

                const number =
                    parseInt(saved, 10);

                if (!isNaN(number)) {
                    position = number;
                }

            }

        } catch (e) {}


        /*
           En vacib hisse:
           sehife yuklense de xetkes
           bir nece defe geri qoyulur.
        */

        sidebar.style.scrollBehavior = "auto";

        sidebar.scrollTop = position;

        requestAnimationFrame(function () {

            sidebar.scrollTop = position;

        });

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 50);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 150);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 300);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 600);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 1000);

    }


    /*
       Xetkes hereket etdikce
       movqeyi yadda saxla.
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            const sidebar =
                getSidebar();

            if (!sidebar) {
                return;
            }

            sidebar.addEventListener(
                "scroll",
                savePosition,
                {
                    passive: true
                }
            );

            restorePosition();

        }
    );


    /*
       PAPKAYA DAXIL OLMADAN EVVEL
       son xetkes yerini yadda saxla.
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            savePosition();

        },
        true
    );


    /*
       Sehife yeniden acilandan sonra
       xetkesi geri getir.
    */

    window.addEventListener(
        "pageshow",
        function () {

            restorePosition();

        }
    );


    /*
       Sehife yuklendi.
    */

    window.addEventListener(
        "load",
        function () {

            restorePosition();

        }
    );


    /*
       Proqram baglananda da yadda saxla.
    */

    window.addEventListener(
        "beforeunload",
        function () {

            savePosition();

        }
    );


})();
</script>
<!-- SCROLL_FIX_START -->
<script>
(function () {

    /*
       ELEKTRON KARGUZARLIQ
       SIDEBAR SCROLL FINAL FIX
    */

    const KEY = "EKARGUZARLIQ_FINAL_SIDEBAR_SCROLL";

    let savedPosition = 0;


    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    /*
       XETKESIN YADDA SAXLA
    */

    function saveScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        savedPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(savedPosition)
            );

        } catch (e) {}

    }


    /*
       YADDA SAXLANMIS XETKESI OXU
    */

    function readScroll() {

        try {

            const value =
                localStorage.getItem(KEY);

            if (value !== null) {

                const number =
                    parseInt(value, 10);

                if (!isNaN(number)) {

                    savedPosition = number;

                }

            }

        } catch (e) {}

    }


    /*
       XETKESI GERI QOY
    */

    function restoreScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        sidebar.scrollTop =
            savedPosition;

    }


    /*
       PAPKAYA KLIK ETMEZDƏN EVVEL
       XETKESI MÜTLƏQ YADDA SAXLA
    */

    document.addEventListener(
        "pointerdown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    document.addEventListener(
        "mousedown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       CLICK-DƏ DƏ YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       XETKESI ELLE HAREKET ETDIRENDƏ
       YENI MOVQE YADDA QALIR
    */

    function connectSidebar() {

        const sidebar =
            getSidebar();

        if (!sidebar) {
            return;
        }

        if (
            sidebar.dataset.finalScrollConnected === "1"
        ) {
            return;
        }

        sidebar.dataset.finalScrollConnected =
            "1";

        sidebar.addEventListener(
            "scroll",
            function () {

                savedPosition =
                    sidebar.scrollTop;

                try {

                    localStorage.setItem(
                        KEY,
                        String(savedPosition)
                    );

                } catch (e) {}

            },
            {
                passive: true
            }
        );

    }


    /*
       SEHIFE YUKLENENDE ESKI MOVQENI OXU
    */

    function startRestore() {

        connectSidebar();

        readScroll();

        restoreScroll();

        /*
           Aktiv menyunun avtomatik olaraq
           xetkesi yuxariya atmasinin qarsisini al.
        */

        let counter = 0;

        const timer =
            setInterval(
                function () {

                    connectSidebar();

                    restoreScroll();

                    counter++;

                    if (counter >= 40) {

                        clearInterval(timer);

                    }

                },
                50
            );

    }


    /*
       DOM HAZIRDIR
    */

    if (
        document.readyState ===
        "loading"
    ) {

        document.addEventListener(
            "DOMContentLoaded",
            startRestore
        );

    } else {

        startRestore();

    }


    /*
       SEHIFE GERI QAYIDANDA
    */

    window.addEventListener(
        "pageshow",
        function () {

            startRestore();

        }
    );


    /*
       SEHIFE BAGLANMAZDAN EVVEL
       SON DEFE YADDA SAXLA
    */

    window.addEventListener(
        "beforeunload",
        function () {

            saveScroll();

        }
    );


})();
</script>
<!-- SCROLL_FIX_END -->

<script>
(function(){
const btn=document.getElementById('mobileMenuBtn'), overlay=document.getElementById('sidebarOverlay');
function closeMenu(){document.body.classList.remove('menu-open')}
if(btn)btn.addEventListener('click',()=>document.body.classList.toggle('menu-open'));
if(overlay)overlay.addEventListener('click',closeMenu);
document.querySelectorAll('.sidebar a').forEach(a=>a.addEventListener('click',()=>{if(innerWidth<=800)closeMenu()}));
document.addEventListener('keydown',e=>{if(e.key==='Escape')closeMenu()});
if('serviceWorker' in navigator)window.addEventListener('load',()=>navigator.serviceWorker.register('/service-worker.js').catch(()=>{}));
})();
</script>

</body>
</html>
"""

INDEX = r"""
{% extends "base" %}

{% block title %}Elektron Kargüzarlıq{% endblock %}

{% block content %}

<div class="cards">

<div class="card">
<div class="card-title">Bütün sənədlər</div>
<div class="card-number">{{ total }}</div>
<div class="card-small">Ümumi sənəd sayı</div>
</div>

<div class="card">
<div class="card-title">İcradadır</div>
<div class="card-number">{{ active }}</div>
<div class="card-small">İcrada olan sənədlər</div>
</div>

<div class="card">
<div class="card-title">Gecikmiş</div>
<div class="card-number">{{ overdue }}</div>
<div class="card-small">Müddəti keçmiş sənədlər</div>
</div>

<div class="card">
<div class="card-title">İcra olunub</div>
<div class="card-number">{{ completed }}</div>
<div class="card-small">Tamamlanmış sənədlər</div>
</div>

</div>

<div class="panel">

<div class="toolbar">

<a class="btn" href="{{ url_for('new_document') }}">
    ➕ Yeni sənəd
</a>

<a class="btn secondary" href="{{ url_for('documents') }}">
    📁 Bütün sənədlər
</a>

</div>

<div class="empty">

<div class="empty-icon">📂</div>

<h2>Elektron Kargüzarlıq Sisteminə xoş gəlmisiniz</h2>

<p>Yeni sənəd əlavə etmək üçün “Yeni sənəd” düyməsinə klikləyin.</p>

</div>

</div>

{% endblock %}
"""

DOCUMENTS = r"""
{% extends "base" %}

{% block title %}Sənədlər{% endblock %}

{% block content %}

{% if current_folder %}

<a class="btn secondary back-btn"
   href="{{ url_for('folders_page') }}">
   ← Əvvələ qayıt
</a>

<h2 class="page-title">
{{ current_folder }} — {{ folder_name }}
</h2>

{% else %}

<h2 class="page-title">Bütün sənədlər</h2>

{% endif %}

<div class="panel">

<div class="toolbar">

<a class="btn" href="{{ url_for('new_document') }}">
    ➕ Yeni sənəd
</a>

<form method="get" style="display:flex;gap:8px;flex:1">

{% if current_folder %}
<input type="hidden" name="folder" value="{{ current_folder }}">
{% endif %}

<input class="search"
       name="q"
       value="{{ q }}"
       placeholder="Axtarış...">

<select name="status" style="width:160px">

<option value="">Bütün statuslar</option>

{% for s in statuses %}

<option value="{{ s }}"
{% if status == s %}selected{% endif %}>
{{ s }}
</option>

{% endfor %}

</select>

<button type="submit">🔍 Axtar</button>

</form>

</div>

<div class="table-wrap">

<table>

<thead>

<tr>
<th>Qeydiyyat №</th>
<th>Papka</th>
<th>Sənəd tarixi</th>
<th>Mövzu</th>
<th>Göndərən</th>
<th>İcraçı</th>
<th>Status</th>
<th>İcra müddəti</th>
<th></th>
</tr>

</thead>

<tbody>

{% for d in docs %}

<tr>

<td>
<a href="{{ url_for('view_document', doc_id=d.id) }}">
<strong>{{ d.reg_no }}</strong>
</a>
</td>

<td>{{ d.folder }}</td>

<td>{{ d.doc_date or '' }}</td>

<td>{{ d.subject or '' }}</td>

<td>{{ d.sender or '' }}</td>

<td>{{ d.executor or '' }}</td>

<td>

{% if d.deadline and d.deadline < today_iso and d.status not in ['İcra olunub','Arxivdə'] %}

<span class="status status-danger">Gecikir</span>

{% elif d.status == 'İcra olunub' %}

<span class="status status-ok">{{ d.status }}</span>

{% else %}

<span class="status">{{ d.status or 'Yeni' }}</span>

{% endif %}

</td>

<td>{{ d.deadline or '' }}</td>

<td>

<a class="btn secondary"
   href="{{ url_for('view_document', doc_id=d.id) }}">
Bax
</a>

</td>

</tr>

{% else %}

<tr>
<td colspan="9">

<div class="empty">

<div class="empty-icon">📂</div>

<h3>Heç bir sənəd yoxdur.</h3>

<p>Yeni sənəd əlavə edə bilərsiniz.</p>

</div>

</td>
</tr>

{% endfor %}

</tbody>

</table>

</div>

</div>

{% endblock %}
"""

FORM = r"""
{% extends "base" %}

{% block title %}{{ title }}{% endblock %}

{% block content %}

<a class="btn secondary back-btn"
   href="{{ url_for('documents', folder=current_folder) if current_folder else url_for('index') }}">
   ← Əvvələ qayıt
</a>

<div class="form-page">

<h2>{{ title }}</h2>

<form method="post" enctype="multipart/form-data">

<div class="form-grid">

<div class="form-group">

<label>Qeydiyyat nömrəsi</label>

<input name="reg_no"
value="{{ document.reg_no if document else reg_no }}"
readonly>

</div>

<div class="form-group">

<label>Papka *</label>

<select name="folder" required>

{% for code,name in folders %}

<option value="{{ code }}"
{% if document and document.folder == code %}selected{% endif %}>
{{ code }} — {{ name }}
</option>

{% endfor %}

</select>

</div>

<div class="form-group">

<label>Sənədin növü</label>

<input name="doc_type"
value="{{ document.doc_type if document else '' }}"
placeholder="Məktub, ərizə, arayış və s.">

</div>

<div class="form-group">

<label>Sənəd tarixi</label>

<input type="date"
name="doc_date"
value="{{ document.doc_date if document else '' }}">

</div>

<div class="form-group">

<label>Daxilolma tarixi</label>

<input type="date"
name="received_date"
value="{{ document.received_date if document else today_iso }}">

</div>

<div class="form-group">

<label>Göndərən</label>

<input name="sender"
value="{{ document.sender if document else '' }}">

</div>

<div class="form-group">

<label>Kimə / ünvanlanan</label>

<input name="recipient"
value="{{ document.recipient if document else '' }}">

</div>

<div class="form-group">

<label>İcraçı</label>

<input name="executor"
value="{{ document.executor if document else '' }}">

</div>

<div class="form-group full">

<label>Mövzu *</label>

<input name="subject"
value="{{ document.subject if document else '' }}"
required>

</div>

<div class="form-group">

<label>İcra müddəti</label>

<input type="date"
name="deadline"
value="{{ document.deadline if document else '' }}">

</div>

<div class="form-group">

<label>Status</label>

<select name="status">

{% for s in statuses %}

<option value="{{ s }}"
{% if document and document.status == s %}selected{% endif %}>
{{ s }}
</option>

{% endfor %}

</select>

</div>

<div class="form-group full">

<label>Dərkənar / Rezolyusiya</label>

<textarea name="resolution"
rows="4">{{ document.resolution if document else '' }}</textarea>

</div>

<div class="form-group full">

<label>Qeyd</label>

<textarea name="notes"
rows="4">{{ document.notes if document else '' }}</textarea>

</div>

<div class="form-group full">

<label>Sənəd faylı</label>

<input type="file" name="file">

{% if document and document.file_name %}

<p>
Mövcud fayl:
<a class="btn"
target="_blank"
href="{{ url_for('preview', filename=document.file_name) }}">
👁 Bax
</a>

<a href="{{ url_for('download', filename=document.file_name) }}">
{{ document.file_name }}
</a>
</p>

{% endif %}

</div>

</div>

<div class="actions">

<button type="submit">
💾 Yadda saxla
</button>

<a class="btn secondary"
href="{{ url_for('documents', folder=current_folder) if current_folder else url_for('index') }}">
Ləğv et
</a>

</div>

</form>

</div>

{% endblock %}
"""

VIEW = r"""
{% extends "base" %}

{% block title %}Sənəd — {{ document.reg_no }}{% endblock %}

{% block content %}

<a class="btn secondary back-btn"
   href="{{ url_for('documents', folder=document.folder) }}">
   ← Əvvələ qayıt
</a>

<div class="form-page">

<div style="display:flex;justify-content:space-between;align-items:center;gap:15px;flex-wrap:wrap">

<h2 style="margin:0">
Sənəd — {{ document.reg_no }}
</h2>

<div>

<a class="btn"
href="{{ url_for('edit_document', doc_id=document.id) }}">
✏ Redaktə
</a>

<form method="post"
action="{{ url_for('delete_document', doc_id=document.id) }}"
style="display:inline"
onsubmit="return confirm('Bu sənədi silmək istəyirsiniz?')">

<button class="btn danger">
🗑 Sil
</button>

</form>

</div>

</div>

<hr>

<div class="detail-grid">

<strong>Qeydiyyat №</strong>
<span>{{ document.reg_no }}</span>

<strong>Papka</strong>
<span>{{ document.folder }} — {{ folder_name }}</span>

<strong>Sənəd növü</strong>
<span>{{ document.doc_type }}</span>

<strong>Sənəd tarixi</strong>
<span>{{ document.doc_date }}</span>

<strong>Daxilolma tarixi</strong>
<span>{{ document.received_date }}</span>

<strong>Mövzu</strong>
<span>{{ document.subject }}</span>

<strong>Göndərən</strong>
<span>{{ document.sender }}</span>

<strong>Kimə</strong>
<span>{{ document.recipient }}</span>

<strong>İcraçı</strong>
<span>{{ document.executor }}</span>

<strong>İcra müddəti</strong>
<span>{{ document.deadline }}</span>

<strong>Status</strong>
<span>{{ document.status }}</span>

<strong>Rezolyusiya</strong>
<span>{{ document.resolution }}</span>

<strong>Qeyd</strong>
<span>{{ document.notes }}</span>

<strong>Yaradılma tarixi</strong>
<span>{{ created_display }}</span>

<strong>Fayl</strong>

<span>

{% if document.file_name %}

<a class="btn"
target="_blank"
href="{{ url_for('preview', filename=document.file_name) }}">
👁 Bax
</a>

<a href="{{ url_for('download', filename=document.file_name) }}">
📄 {{ document.file_name }}
</a>

{% else %}

Fayl əlavə edilməyib.

{% endif %}

</span>

</div>

</div>

{% endblock %}
"""

FOLDERS_PAGE = r"""
{% extends "base" %}

{% block title %}Papkalar{% endblock %}

{% block content %}

<a class="btn secondary back-btn"
   href="{{ url_for('folders_page') }}">
   ← Əvvələ qayıt
</a>

<h2 class="page-title">Prorektorluğun papka sistemi</h2>

<div class="folder-list">

{% for code,name in folders %}

<a class="folder-card"
href="{{ url_for('documents', folder=code) }}">

<div style="font-size:24px">
{{ icons.get(code,'📁') }}
</div>

<strong>{{ code }}</strong>

<br>

{{ name }}

</a>

{% endfor %}

</div>

{% endblock %}
"""

# ============================================================
# TEMPLATE RENDER
# ============================================================

from jinja2 import DictLoader

app.jinja_loader = DictLoader({
    "base": HTML,
    "index.html": INDEX,
    "documents.html": DOCUMENTS,
    "form.html": FORM,
    "view.html": VIEW,
    "folders.html": FOLDERS_PAGE,
})


def render(template, **kwargs):

    kwargs.setdefault("folders", FOLDERS)
    kwargs.setdefault("statuses", STATUSES)
    kwargs.setdefault("icons", ICONS)

    now = datetime.now()

    kwargs.setdefault(
        "today_iso",
        now.strftime("%Y-%m-%d")
    )

    kwargs.setdefault(
        "now_display",
        now.strftime("%d.%m.%Y %H:%M")
    )

    return app.jinja_env.get_template(template).render(**kwargs)


# ============================================================
# ANA SƏHİFƏ
# ============================================================

@app.route("/")
def index():

    conn = db()

    total = conn.execute(
        "SELECT COUNT(*) AS c FROM documents"
    ).fetchone()["c"]

    active = conn.execute(
        "SELECT COUNT(*) AS c FROM documents WHERE status='İcradadır'"
    ).fetchone()["c"]

    completed = conn.execute(
        "SELECT COUNT(*) AS c FROM documents WHERE status='İcra olunub'"
    ).fetchone()["c"]

    overdue = conn.execute("""
        SELECT COUNT(*) AS c
        FROM documents
        WHERE deadline IS NOT NULL
        AND deadline != ''
        AND deadline < ?
        AND status NOT IN ('İcra olunub','Arxivdə')
    """, (date.today().isoformat(),)).fetchone()["c"]

    conn.close()

    return render(
        "index.html",
        total=total,
        active=active,
        completed=completed,
        overdue=overdue,
        current_folder=None
    )


# ============================================================
# SƏNƏDLƏR
# ============================================================

@app.route("/documents")
def documents():

    folder = request.args.get("folder", "").strip()
    q = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()

    conn = db()

    sql = "SELECT * FROM documents WHERE 1=1"
    params = []

    if folder:
        sql += " AND folder = ?"
        params.append(folder)

    if q:
        sql += """
        AND (
            reg_no LIKE ?
            OR subject LIKE ?
            OR sender LIKE ?
            OR executor LIKE ?
            OR recipient LIKE ?
        )
        """

        value = f"%{q}%"

        params.extend([
            value,
            value,
            value,
            value,
            value
        ])

    if status:
        sql += " AND status = ?"
        params.append(status)

    sql += " ORDER BY id DESC"

    docs = conn.execute(sql, params).fetchall()

    conn.close()

    return render(
        "documents.html",
        docs=docs,
        q=q,
        status=status,
        current_folder=folder or None,
        folder_name=FOLDER_NAMES.get(folder, ""),
    )


# ============================================================
# YENİ SƏNƏD
# ============================================================

@app.route("/new", methods=["GET", "POST"])
def new_document():

    reg_no = next_reg_no()

    if request.method == "POST":

        file = request.files.get("file")

        filename = ""

        if file and file.filename:

            original = secure_filename(file.filename)

            stamp = datetime.now().strftime(
                "%Y%m%d%H%M%S"
            )

            filename = stamp + "_" + original

            file.save(
                os.path.join(
                    UPLOAD_DIR,
                    filename
                )
            )

        form = request.form

        conn = db()

        conn.execute("""
            INSERT INTO documents (
                reg_no,
                folder,
                doc_type,
                doc_date,
                received_date,
                subject,
                sender,
                recipient,
                executor,
                deadline,
                status,
                resolution,
                notes,
                file_name,
                created_at
            )
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            form.get("reg_no") or reg_no,
            form.get("folder") or "01",
            normalize(form.get("doc_type")),
            normalize(form.get("doc_date")),
            normalize(form.get("received_date")),
            normalize(form.get("subject")),
            normalize(form.get("sender")),
            normalize(form.get("recipient")),
            normalize(form.get("executor")),
            normalize(form.get("deadline")),
            form.get("status") or "Yeni",
            normalize(form.get("resolution")),
            normalize(form.get("notes")),
            filename,
            datetime.now().isoformat(timespec="seconds")
        ))

        conn.commit()
        conn.close()

        return redirect(
            url_for(
                "documents",
                folder=form.get("folder") or "01"
            )
        )

    return render(
        "form.html",
        title="Yeni sənəd",
        document=None,
        reg_no=reg_no,
        current_folder=None
    )


# ============================================================
# SƏNƏDƏ BAX
# ============================================================

@app.route("/document/<int:doc_id>")
def view_document(doc_id):

    conn = db()

    document = conn.execute(
        "SELECT * FROM documents WHERE id=?",
        (doc_id,)
    ).fetchone()

    conn.close()

    if not document:
        return "Sənəd tapılmadı", 404

    return render(
        "view.html",
        document=document,
        current_folder=document["folder"],
        folder_name=FOLDER_NAMES.get(
            document["folder"],
            ""
        ),
        created_display=format_datetime(
            document["created_at"]
        )
    )


# ============================================================
# REDAKTƏ
# ============================================================

@app.route("/document/<int:doc_id>/edit", methods=["GET", "POST"])
def edit_document(doc_id):

    conn = db()

    document = conn.execute(
        "SELECT * FROM documents WHERE id=?",
        (doc_id,)
    ).fetchone()

    conn.close()

    if not document:
        return "Sənəd tapılmadı", 404

    if request.method == "POST":

        file = request.files.get("file")

        filename = document["file_name"] or ""

        if file and file.filename:

            original = secure_filename(
                file.filename
            )

            stamp = datetime.now().strftime(
                "%Y%m%d%H%M%S"
            )

            filename = stamp + "_" + original

            file.save(
                os.path.join(
                    UPLOAD_DIR,
                    filename
                )
            )

        form = request.form

        conn = db()

        conn.execute("""
            UPDATE documents SET
                folder=?,
                doc_type=?,
                doc_date=?,
                received_date=?,
                subject=?,
                sender=?,
                recipient=?,
                executor=?,
                deadline=?,
                status=?,
                resolution=?,
                notes=?,
                file_name=?
            WHERE id=?
        """, (
            form.get("folder") or document["folder"],
            normalize(form.get("doc_type")),
            normalize(form.get("doc_date")),
            normalize(form.get("received_date")),
            normalize(form.get("subject")),
            normalize(form.get("sender")),
            normalize(form.get("recipient")),
            normalize(form.get("executor")),
            normalize(form.get("deadline")),
            form.get("status") or "Yeni",
            normalize(form.get("resolution")),
            normalize(form.get("notes")),
            filename,
            doc_id
        ))

        conn.commit()
        conn.close()

        return redirect(
            url_for(
                "view_document",
                doc_id=doc_id
            )
        )

    return render(
        "form.html",
        title="Sənədi redaktə et",
        document=document,
        reg_no=document["reg_no"],
        current_folder=document["folder"]
    )


# ============================================================
# SİL
# ============================================================

@app.route("/document/<int:doc_id>/delete", methods=["POST"])
def delete_document(doc_id):

    conn = db()

    document = conn.execute(
        "SELECT * FROM documents WHERE id=?",
        (doc_id,)
    ).fetchone()

    if document:

        filename = document["file_name"]

        if filename:

            file_path = os.path.join(
                UPLOAD_DIR,
                filename
            )

            if os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except Exception:
                    pass

        conn.execute(
            "DELETE FROM documents WHERE id=?",
            (doc_id,)
        )

        conn.commit()

    conn.close()

    return redirect(
        url_for(
            "documents",
            folder=document["folder"] if document else None
        )
    )


# ============================================================
# PREVIEW
# ============================================================

@app.route("/preview/<path:filename>")
def preview(filename):

    safe_name = os.path.basename(filename)

    file_path = os.path.join(
        UPLOAD_DIR,
        safe_name
    )

    if not os.path.isfile(file_path):
        return "Fayl tapılmadı", 404

    ext = Path(file_path).suffix.lower()

    # Excel
    if ext in [".xlsx",".xlsm",".xltx",".xltm"]:

        try:

            from openpyxl import load_workbook

            wb = load_workbook(
                file_path,
                read_only=True,
                data_only=True
            )

            sheets = []

            for ws in wb.worksheets:

                rows_html = []

                for row in ws.iter_rows(values_only=True):

                    cells = []

                    for value in row:
                        cells.append(
                            "<td>" +
                            str(value or "") +
                            "</td>"
                        )

                    rows_html.append(
                        "<tr>" +
                        "".join(cells) +
                        "</tr>"
                    )

                sheets.append(
                    "<h2>" +
                    ws.title +
                    "</h2>" +
                    "<div class='table-box'>" +
                    "<table>" +
                    "".join(rows_html) +
                    "</table>" +
                    "</div>"
                )

            wb.close()

            html = """
<!DOCTYPE html>
<html lang="az">
<head>
<meta charset="UTF-8">
<title>Sənədə baxış</title>
<style>
body {
    margin:0;
    padding:30px;
    background:#f3f6fa;
    font-family:Arial,sans-serif;
    color:#173a64;
}
.box {
    background:white;
    padding:25px;
    border-radius:12px;
}
.table-box {
    overflow:auto;
    background:white;
    margin-bottom:20px;
}
table {
    border-collapse:collapse;
    width:100%;
}
td {
    border:1px solid #d7e0ea;
    padding:8px;
}
</style>
</head>
<body>
<div class="box">
<h1>📊 Sənədə baxış</h1>
<p><b>Fayl:</b> __FILENAME__</p>
__SHEETS__
</div>


<script>
(function () {

    const MAIN_SCROLL_KEY =
        "ekarguzarlıq_main_page_scroll";

    function saveMainScroll() {

        try {

            sessionStorage.setItem(
                MAIN_SCROLL_KEY,
                String(window.scrollY || window.pageYOffset || 0)
            );

        } catch (e) {}
    }


    function restoreMainScroll() {

        try {

            const saved =
                sessionStorage.getItem(
                    MAIN_SCROLL_KEY
                );

            if (saved !== null) {

                const position =
                    parseInt(saved, 10) || 0;

                window.scrollTo(
                    0,
                    position
                );

            }

        } catch (e) {}
    }


    /*
       ESAS SEHIFE XETKESI HAREKET ETDIKCE
       YADDA SAXLA
    */

    window.addEventListener(
        "scroll",
        saveMainScroll,
        { passive: true }
    );


    /*
       ISTENILEN FAYLA/SENEDE KLIK EDILMEZDƏN
       EVVEL XETKESI YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest("a");

            if (!link) {
                return;
            }

            if (
                link.target === "_blank" ||
                link.hasAttribute("download")
            ) {
                return;
            }

            saveMainScroll();

        },
        true
    );


    /*
       SEHIFE GERI ACILANDA
       EVVELKI XETKESI QAYTAR
    */

    window.addEventListener(
        "pageshow",
        function () {

            restoreMainScroll();

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                150
            );

            setTimeout(
                restoreMainScroll,
                300
            );

            setTimeout(
                restoreMainScroll,
                600
            );

        }
    );


    /*
       SEHIFE YUKLENENDE
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                200
            );

        }
    );


    /*
       PROQRAM BAGLANANDA
    */

    window.addEventListener(
        "beforeunload",
        saveMainScroll
    );

})();
</script>
<script>
(function () {

    const KEY = "EKARGUZARLIQ_SIDEBAR_SCROLL";

    let lastPosition = 0;

    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    function savePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        lastPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(lastPosition)
            );

        } catch (e) {}

    }


    function restorePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        let position = lastPosition;

        try {

            const saved =
                localStorage.getItem(KEY);

            if (saved !== null) {

                const number =
                    parseInt(saved, 10);

                if (!isNaN(number)) {
                    position = number;
                }

            }

        } catch (e) {}


        /*
           En vacib hisse:
           sehife yuklense de xetkes
           bir nece defe geri qoyulur.
        */

        sidebar.style.scrollBehavior = "auto";

        sidebar.scrollTop = position;

        requestAnimationFrame(function () {

            sidebar.scrollTop = position;

        });

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 50);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 150);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 300);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 600);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 1000);

    }


    /*
       Xetkes hereket etdikce
       movqeyi yadda saxla.
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            const sidebar =
                getSidebar();

            if (!sidebar) {
                return;
            }

            sidebar.addEventListener(
                "scroll",
                savePosition,
                {
                    passive: true
                }
            );

            restorePosition();

        }
    );


    /*
       PAPKAYA DAXIL OLMADAN EVVEL
       son xetkes yerini yadda saxla.
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            savePosition();

        },
        true
    );


    /*
       Sehife yeniden acilandan sonra
       xetkesi geri getir.
    */

    window.addEventListener(
        "pageshow",
        function () {

            restorePosition();

        }
    );


    /*
       Sehife yuklendi.
    */

    window.addEventListener(
        "load",
        function () {

            restorePosition();

        }
    );


    /*
       Proqram baglananda da yadda saxla.
    */

    window.addEventListener(
        "beforeunload",
        function () {

            savePosition();

        }
    );


})();
</script>
<!-- SCROLL_FIX_START -->
<script>
(function () {

    /*
       ELEKTRON KARGUZARLIQ
       SIDEBAR SCROLL FINAL FIX
    */

    const KEY = "EKARGUZARLIQ_FINAL_SIDEBAR_SCROLL";

    let savedPosition = 0;


    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    /*
       XETKESIN YADDA SAXLA
    */

    function saveScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        savedPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(savedPosition)
            );

        } catch (e) {}

    }


    /*
       YADDA SAXLANMIS XETKESI OXU
    */

    function readScroll() {

        try {

            const value =
                localStorage.getItem(KEY);

            if (value !== null) {

                const number =
                    parseInt(value, 10);

                if (!isNaN(number)) {

                    savedPosition = number;

                }

            }

        } catch (e) {}

    }


    /*
       XETKESI GERI QOY
    */

    function restoreScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        sidebar.scrollTop =
            savedPosition;

    }


    /*
       PAPKAYA KLIK ETMEZDƏN EVVEL
       XETKESI MÜTLƏQ YADDA SAXLA
    */

    document.addEventListener(
        "pointerdown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    document.addEventListener(
        "mousedown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       CLICK-DƏ DƏ YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       XETKESI ELLE HAREKET ETDIRENDƏ
       YENI MOVQE YADDA QALIR
    */

    function connectSidebar() {

        const sidebar =
            getSidebar();

        if (!sidebar) {
            return;
        }

        if (
            sidebar.dataset.finalScrollConnected === "1"
        ) {
            return;
        }

        sidebar.dataset.finalScrollConnected =
            "1";

        sidebar.addEventListener(
            "scroll",
            function () {

                savedPosition =
                    sidebar.scrollTop;

                try {

                    localStorage.setItem(
                        KEY,
                        String(savedPosition)
                    );

                } catch (e) {}

            },
            {
                passive: true
            }
        );

    }


    /*
       SEHIFE YUKLENENDE ESKI MOVQENI OXU
    */

    function startRestore() {

        connectSidebar();

        readScroll();

        restoreScroll();

        /*
           Aktiv menyunun avtomatik olaraq
           xetkesi yuxariya atmasinin qarsisini al.
        */

        let counter = 0;

        const timer =
            setInterval(
                function () {

                    connectSidebar();

                    restoreScroll();

                    counter++;

                    if (counter >= 40) {

                        clearInterval(timer);

                    }

                },
                50
            );

    }


    /*
       DOM HAZIRDIR
    */

    if (
        document.readyState ===
        "loading"
    ) {

        document.addEventListener(
            "DOMContentLoaded",
            startRestore
        );

    } else {

        startRestore();

    }


    /*
       SEHIFE GERI QAYIDANDA
    */

    window.addEventListener(
        "pageshow",
        function () {

            startRestore();

        }
    );


    /*
       SEHIFE BAGLANMAZDAN EVVEL
       SON DEFE YADDA SAXLA
    */

    window.addEventListener(
        "beforeunload",
        function () {

            saveScroll();

        }
    );


})();
</script>
<!-- SCROLL_FIX_END -->
</body>
</html>
"""

            html = html.replace(
                "__FILENAME__",
                safe_name
            )

            html = html.replace(
                "__SHEETS__",
                "".join(sheets)
            )

            return Response(
                html,
                mimetype="text/html"
            )

        except Exception as e:

            return Response(
                "<h2>Excel faylı açıla bilmədi</h2><p>" +
                str(e) +
                "</p>",
                status=500,
                mimetype="text/html"
            )

    # Word
    if ext == ".docx":

        try:

            from docx import Document

            doc = Document(file_path)

            content = []

            for paragraph in doc.paragraphs:

                text = paragraph.text.strip()

                if text:
                    content.append(
                        "<p>" +
                        text +
                        "</p>"
                    )

            for table in doc.tables:

                rows = []

                for row in table.rows:

                    cells = []

                    for cell in row.cells:

                        cells.append(
                            "<td>" +
                            cell.text +
                            "</td>"
                        )

                    rows.append(
                        "<tr>" +
                        "".join(cells) +
                        "</tr>"
                    )

                content.append(
                    "<table>" +
                    "".join(rows) +
                    "</table>"
                )

            html = """
<!DOCTYPE html>
<html lang="az">
<head>
<meta charset="UTF-8">
<title>Sənədə baxış</title>
<style>
body {
    margin:0;
    padding:30px;
    background:#f3f6fa;
    font-family:Arial,sans-serif;
}
.box {
    max-width:1100px;
    margin:auto;
    background:white;
    padding:30px;
    border-radius:12px;
}
table {
    border-collapse:collapse;
    width:100%;
    margin-top:20px;
}
td {
    border:1px solid #d7e0ea;
    padding:9px;
}
p {
    line-height:1.7;
}
</style>
</head>
<body>
<div class="box">
<h1>📄 Sənədə baxış</h1>
<p><b>Fayl:</b> __FILENAME__</p>
<hr>
__CONTENT__
</div>


<script>
(function () {

    const MAIN_SCROLL_KEY =
        "ekarguzarlıq_main_page_scroll";

    function saveMainScroll() {

        try {

            sessionStorage.setItem(
                MAIN_SCROLL_KEY,
                String(window.scrollY || window.pageYOffset || 0)
            );

        } catch (e) {}
    }


    function restoreMainScroll() {

        try {

            const saved =
                sessionStorage.getItem(
                    MAIN_SCROLL_KEY
                );

            if (saved !== null) {

                const position =
                    parseInt(saved, 10) || 0;

                window.scrollTo(
                    0,
                    position
                );

            }

        } catch (e) {}
    }


    /*
       ESAS SEHIFE XETKESI HAREKET ETDIKCE
       YADDA SAXLA
    */

    window.addEventListener(
        "scroll",
        saveMainScroll,
        { passive: true }
    );


    /*
       ISTENILEN FAYLA/SENEDE KLIK EDILMEZDƏN
       EVVEL XETKESI YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest("a");

            if (!link) {
                return;
            }

            if (
                link.target === "_blank" ||
                link.hasAttribute("download")
            ) {
                return;
            }

            saveMainScroll();

        },
        true
    );


    /*
       SEHIFE GERI ACILANDA
       EVVELKI XETKESI QAYTAR
    */

    window.addEventListener(
        "pageshow",
        function () {

            restoreMainScroll();

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                150
            );

            setTimeout(
                restoreMainScroll,
                300
            );

            setTimeout(
                restoreMainScroll,
                600
            );

        }
    );


    /*
       SEHIFE YUKLENENDE
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                200
            );

        }
    );


    /*
       PROQRAM BAGLANANDA
    */

    window.addEventListener(
        "beforeunload",
        saveMainScroll
    );

})();
</script>
<script>
(function () {

    const KEY = "EKARGUZARLIQ_SIDEBAR_SCROLL";

    let lastPosition = 0;

    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    function savePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        lastPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(lastPosition)
            );

        } catch (e) {}

    }


    function restorePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        let position = lastPosition;

        try {

            const saved =
                localStorage.getItem(KEY);

            if (saved !== null) {

                const number =
                    parseInt(saved, 10);

                if (!isNaN(number)) {
                    position = number;
                }

            }

        } catch (e) {}


        /*
           En vacib hisse:
           sehife yuklense de xetkes
           bir nece defe geri qoyulur.
        */

        sidebar.style.scrollBehavior = "auto";

        sidebar.scrollTop = position;

        requestAnimationFrame(function () {

            sidebar.scrollTop = position;

        });

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 50);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 150);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 300);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 600);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 1000);

    }


    /*
       Xetkes hereket etdikce
       movqeyi yadda saxla.
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            const sidebar =
                getSidebar();

            if (!sidebar) {
                return;
            }

            sidebar.addEventListener(
                "scroll",
                savePosition,
                {
                    passive: true
                }
            );

            restorePosition();

        }
    );


    /*
       PAPKAYA DAXIL OLMADAN EVVEL
       son xetkes yerini yadda saxla.
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            savePosition();

        },
        true
    );


    /*
       Sehife yeniden acilandan sonra
       xetkesi geri getir.
    */

    window.addEventListener(
        "pageshow",
        function () {

            restorePosition();

        }
    );


    /*
       Sehife yuklendi.
    */

    window.addEventListener(
        "load",
        function () {

            restorePosition();

        }
    );


    /*
       Proqram baglananda da yadda saxla.
    */

    window.addEventListener(
        "beforeunload",
        function () {

            savePosition();

        }
    );


})();
</script>
<!-- SCROLL_FIX_START -->
<script>
(function () {

    /*
       ELEKTRON KARGUZARLIQ
       SIDEBAR SCROLL FINAL FIX
    */

    const KEY = "EKARGUZARLIQ_FINAL_SIDEBAR_SCROLL";

    let savedPosition = 0;


    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    /*
       XETKESIN YADDA SAXLA
    */

    function saveScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        savedPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(savedPosition)
            );

        } catch (e) {}

    }


    /*
       YADDA SAXLANMIS XETKESI OXU
    */

    function readScroll() {

        try {

            const value =
                localStorage.getItem(KEY);

            if (value !== null) {

                const number =
                    parseInt(value, 10);

                if (!isNaN(number)) {

                    savedPosition = number;

                }

            }

        } catch (e) {}

    }


    /*
       XETKESI GERI QOY
    */

    function restoreScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        sidebar.scrollTop =
            savedPosition;

    }


    /*
       PAPKAYA KLIK ETMEZDƏN EVVEL
       XETKESI MÜTLƏQ YADDA SAXLA
    */

    document.addEventListener(
        "pointerdown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    document.addEventListener(
        "mousedown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       CLICK-DƏ DƏ YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       XETKESI ELLE HAREKET ETDIRENDƏ
       YENI MOVQE YADDA QALIR
    */

    function connectSidebar() {

        const sidebar =
            getSidebar();

        if (!sidebar) {
            return;
        }

        if (
            sidebar.dataset.finalScrollConnected === "1"
        ) {
            return;
        }

        sidebar.dataset.finalScrollConnected =
            "1";

        sidebar.addEventListener(
            "scroll",
            function () {

                savedPosition =
                    sidebar.scrollTop;

                try {

                    localStorage.setItem(
                        KEY,
                        String(savedPosition)
                    );

                } catch (e) {}

            },
            {
                passive: true
            }
        );

    }


    /*
       SEHIFE YUKLENENDE ESKI MOVQENI OXU
    */

    function startRestore() {

        connectSidebar();

        readScroll();

        restoreScroll();

        /*
           Aktiv menyunun avtomatik olaraq
           xetkesi yuxariya atmasinin qarsisini al.
        */

        let counter = 0;

        const timer =
            setInterval(
                function () {

                    connectSidebar();

                    restoreScroll();

                    counter++;

                    if (counter >= 40) {

                        clearInterval(timer);

                    }

                },
                50
            );

    }


    /*
       DOM HAZIRDIR
    */

    if (
        document.readyState ===
        "loading"
    ) {

        document.addEventListener(
            "DOMContentLoaded",
            startRestore
        );

    } else {

        startRestore();

    }


    /*
       SEHIFE GERI QAYIDANDA
    */

    window.addEventListener(
        "pageshow",
        function () {

            startRestore();

        }
    );


    /*
       SEHIFE BAGLANMAZDAN EVVEL
       SON DEFE YADDA SAXLA
    */

    window.addEventListener(
        "beforeunload",
        function () {

            saveScroll();

        }
    );


})();
</script>
<!-- SCROLL_FIX_END -->
</body>
</html>
"""

            html = html.replace(
                "__FILENAME__",
                safe_name
            )

            html = html.replace(
                "__CONTENT__",
                "".join(content)
            )

            return Response(
                html,
                mimetype="text/html"
            )

        except Exception as e:

            return Response(
                "<h2>Word faylı açıla bilmədi</h2><p>" +
                str(e) +
                "</p>",
                status=500,
                mimetype="text/html"
            )

    # TXT / CSV / LOG
    if ext in [".txt",".csv",".log"]:

        try:

            content = Path(
                file_path
            ).read_text(
                encoding="utf-8",
                errors="replace"
            )

            html = """
<!DOCTYPE html>
<html lang="az">
<head>
<meta charset="UTF-8">
<title>Sənədə baxış</title>
<style>
body {
    margin:0;
    padding:30px;
    background:#f3f6fa;
    font-family:Arial,sans-serif;
}
.box {
    max-width:1100px;
    margin:auto;
    background:white;
    padding:30px;
    border-radius:12px;
    white-space:pre-wrap;
}
</style>
</head>
<body>
<div class="box">
<h2>📄 __FILENAME__</h2>
<hr>
__CONTENT__
</div>


<script>
(function () {

    const MAIN_SCROLL_KEY =
        "ekarguzarlıq_main_page_scroll";

    function saveMainScroll() {

        try {

            sessionStorage.setItem(
                MAIN_SCROLL_KEY,
                String(window.scrollY || window.pageYOffset || 0)
            );

        } catch (e) {}
    }


    function restoreMainScroll() {

        try {

            const saved =
                sessionStorage.getItem(
                    MAIN_SCROLL_KEY
                );

            if (saved !== null) {

                const position =
                    parseInt(saved, 10) || 0;

                window.scrollTo(
                    0,
                    position
                );

            }

        } catch (e) {}
    }


    /*
       ESAS SEHIFE XETKESI HAREKET ETDIKCE
       YADDA SAXLA
    */

    window.addEventListener(
        "scroll",
        saveMainScroll,
        { passive: true }
    );


    /*
       ISTENILEN FAYLA/SENEDE KLIK EDILMEZDƏN
       EVVEL XETKESI YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest("a");

            if (!link) {
                return;
            }

            if (
                link.target === "_blank" ||
                link.hasAttribute("download")
            ) {
                return;
            }

            saveMainScroll();

        },
        true
    );


    /*
       SEHIFE GERI ACILANDA
       EVVELKI XETKESI QAYTAR
    */

    window.addEventListener(
        "pageshow",
        function () {

            restoreMainScroll();

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                150
            );

            setTimeout(
                restoreMainScroll,
                300
            );

            setTimeout(
                restoreMainScroll,
                600
            );

        }
    );


    /*
       SEHIFE YUKLENENDE
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            setTimeout(
                restoreMainScroll,
                50
            );

            setTimeout(
                restoreMainScroll,
                200
            );

        }
    );


    /*
       PROQRAM BAGLANANDA
    */

    window.addEventListener(
        "beforeunload",
        saveMainScroll
    );

})();
</script>
<script>
(function () {

    const KEY = "EKARGUZARLIQ_SIDEBAR_SCROLL";

    let lastPosition = 0;

    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    function savePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        lastPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(lastPosition)
            );

        } catch (e) {}

    }


    function restorePosition() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        let position = lastPosition;

        try {

            const saved =
                localStorage.getItem(KEY);

            if (saved !== null) {

                const number =
                    parseInt(saved, 10);

                if (!isNaN(number)) {
                    position = number;
                }

            }

        } catch (e) {}


        /*
           En vacib hisse:
           sehife yuklense de xetkes
           bir nece defe geri qoyulur.
        */

        sidebar.style.scrollBehavior = "auto";

        sidebar.scrollTop = position;

        requestAnimationFrame(function () {

            sidebar.scrollTop = position;

        });

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 50);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 150);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 300);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 600);

        setTimeout(function () {

            sidebar.scrollTop = position;

        }, 1000);

    }


    /*
       Xetkes hereket etdikce
       movqeyi yadda saxla.
    */

    document.addEventListener(
        "DOMContentLoaded",
        function () {

            const sidebar =
                getSidebar();

            if (!sidebar) {
                return;
            }

            sidebar.addEventListener(
                "scroll",
                savePosition,
                {
                    passive: true
                }
            );

            restorePosition();

        }
    );


    /*
       PAPKAYA DAXIL OLMADAN EVVEL
       son xetkes yerini yadda saxla.
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            savePosition();

        },
        true
    );


    /*
       Sehife yeniden acilandan sonra
       xetkesi geri getir.
    */

    window.addEventListener(
        "pageshow",
        function () {

            restorePosition();

        }
    );


    /*
       Sehife yuklendi.
    */

    window.addEventListener(
        "load",
        function () {

            restorePosition();

        }
    );


    /*
       Proqram baglananda da yadda saxla.
    */

    window.addEventListener(
        "beforeunload",
        function () {

            savePosition();

        }
    );


})();
</script>
<!-- SCROLL_FIX_START -->
<script>
(function () {

    /*
       ELEKTRON KARGUZARLIQ
       SIDEBAR SCROLL FINAL FIX
    */

    const KEY = "EKARGUZARLIQ_FINAL_SIDEBAR_SCROLL";

    let savedPosition = 0;


    function getSidebar() {

        return document.querySelector(".sidebar");

    }


    /*
       XETKESIN YADDA SAXLA
    */

    function saveScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        savedPosition =
            sidebar.scrollTop;

        try {

            localStorage.setItem(
                KEY,
                String(savedPosition)
            );

        } catch (e) {}

    }


    /*
       YADDA SAXLANMIS XETKESI OXU
    */

    function readScroll() {

        try {

            const value =
                localStorage.getItem(KEY);

            if (value !== null) {

                const number =
                    parseInt(value, 10);

                if (!isNaN(number)) {

                    savedPosition = number;

                }

            }

        } catch (e) {}

    }


    /*
       XETKESI GERI QOY
    */

    function restoreScroll() {

        const sidebar = getSidebar();

        if (!sidebar) {
            return;
        }

        sidebar.scrollTop =
            savedPosition;

    }


    /*
       PAPKAYA KLIK ETMEZDƏN EVVEL
       XETKESI MÜTLƏQ YADDA SAXLA
    */

    document.addEventListener(
        "pointerdown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    document.addEventListener(
        "mousedown",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       CLICK-DƏ DƏ YADDA SAXLA
    */

    document.addEventListener(
        "click",
        function (event) {

            const link =
                event.target.closest(
                    "a.menu-item"
                );

            if (!link) {
                return;
            }

            saveScroll();

        },
        true
    );


    /*
       XETKESI ELLE HAREKET ETDIRENDƏ
       YENI MOVQE YADDA QALIR
    */

    function connectSidebar() {

        const sidebar =
            getSidebar();

        if (!sidebar) {
            return;
        }

        if (
            sidebar.dataset.finalScrollConnected === "1"
        ) {
            return;
        }

        sidebar.dataset.finalScrollConnected =
            "1";

        sidebar.addEventListener(
            "scroll",
            function () {

                savedPosition =
                    sidebar.scrollTop;

                try {

                    localStorage.setItem(
                        KEY,
                        String(savedPosition)
                    );

                } catch (e) {}

            },
            {
                passive: true
            }
        );

    }


    /*
       SEHIFE YUKLENENDE ESKI MOVQENI OXU
    */

    function startRestore() {

        connectSidebar();

        readScroll();

        restoreScroll();

        /*
           Aktiv menyunun avtomatik olaraq
           xetkesi yuxariya atmasinin qarsisini al.
        */

        let counter = 0;

        const timer =
            setInterval(
                function () {

                    connectSidebar();

                    restoreScroll();

                    counter++;

                    if (counter >= 40) {

                        clearInterval(timer);

                    }

                },
                50
            );

    }


    /*
       DOM HAZIRDIR
    */

    if (
        document.readyState ===
        "loading"
    ) {

        document.addEventListener(
            "DOMContentLoaded",
            startRestore
        );

    } else {

        startRestore();

    }


    /*
       SEHIFE GERI QAYIDANDA
    */

    window.addEventListener(
        "pageshow",
        function () {

            startRestore();

        }
    );


    /*
       SEHIFE BAGLANMAZDAN EVVEL
       SON DEFE YADDA SAXLA
    */

    window.addEventListener(
        "beforeunload",
        function () {

            saveScroll();

        }
    );


})();
</script>
<!-- SCROLL_FIX_END -->
</body>
</html>
"""

            html = html.replace(
                "__FILENAME__",
                safe_name
            )

            html = html.replace(
                "__CONTENT__",
                content
            )

            return Response(
                html,
                mimetype="text/html"
            )

        except Exception:
            pass

    mime = mimetypes.guess_type(
        file_path
    )[0] or "application/octet-stream"

    return send_from_directory(
        UPLOAD_DIR,
        safe_name,
        as_attachment=False,
        mimetype=mime
    )


# ============================================================
# DOWNLOAD
# ============================================================

@app.route("/download/<path:filename>")
def download(filename):

    return send_from_directory(
        UPLOAD_DIR,
        os.path.basename(filename),
        as_attachment=True
    )


# ============================================================
# PAPKALAR
# ============================================================

@app.route("/folders")
def folders_page():

    return render(
        "folders.html",
        current_folder=None
    )


# ============================================================
# PWA / TELEFONA QURAŞDIRMA
# ============================================================

@app.route("/manifest.webmanifest")
def manifest():
    return jsonify({
        "name": "Elektron Kargüzarlıq",
        "short_name": "E-Kargüzarlıq",
        "description": "Sənəd uçotu və icra nəzarəti sistemi",
        "start_url": "/",
        "scope": "/",
        "display": "standalone",
        "background_color": "#f4f7fb",
        "theme_color": "#062b52",
        "lang": "az",
        "icons": [{"src": "/app-icon.svg", "sizes": "any", "type": "image/svg+xml", "purpose": "any maskable"}]
    })

@app.route("/app-icon.svg")
def app_icon():
    svg = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 512 512"><rect width="512" height="512" rx="96" fill="#062b52"/><path d="M128 96h192l64 64v256H128z" fill="#fff"/><path d="M320 96v80h64" fill="#cfe4ff"/><path d="M176 240h160M176 288h160M176 336h120" stroke="#1473e6" stroke-width="24" stroke-linecap="round"/></svg>"""
    return Response(svg, mimetype="image/svg+xml")

@app.route("/service-worker.js")
def service_worker():
    js = """self.addEventListener('install',e=>self.skipWaiting());self.addEventListener('activate',e=>e.waitUntil(self.clients.claim()));self.addEventListener('fetch',e=>{});"""
    return Response(js, mimetype="application/javascript", headers={"Cache-Control":"no-cache"})

@app.route("/health")
def health():
    return {"status": "ok"}

# ============================================================
# BAŞLAT
# ============================================================

def open_browser():

    webbrowser.open(
        "http://127.0.0.1:5000"
    )


if __name__ == "__main__":

    init_db()

    print("=" * 60)
    print(" ELEKTRON KARGÜZARLIQ")
    print("=" * 60)
    print("DB:", DB)
    print("UPLOADS:", UPLOAD_DIR)
    port = int(os.environ.get("PORT", "5000"))
    host = os.environ.get("HOST", "0.0.0.0")
    print(f"http://{host}:{port}")
    print("=" * 60)

    threading.Timer(
        1.2,
        open_browser
    ).start()

    app.run(
        host=host,
        port=port,
        debug=False,
        use_reloader=False
    )










